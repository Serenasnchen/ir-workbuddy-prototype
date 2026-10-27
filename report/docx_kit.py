"""Small python-docx toolkit for a clean Chinese business report."""
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SANS = "Microsoft YaHei"
SERIF = "SimSun"
INK = RGBColor(0x1B, 0x24, 0x33)
MUTED = RGBColor(0x5B, 0x65, 0x77)
BLUE = RGBColor(0x24, 0x59, 0xD6)
AMBER = RGBColor(0xB9, 0x7E, 0x12)
RED = RGBColor(0xB8, 0x43, 0x2F)


def _font(run, name=SANS, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _style_font(style, name, size, bold=False, color=INK):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    st = doc.styles
    _style_font(st["Normal"], SANS, 10)
    pf = st["Normal"].paragraph_format
    pf.line_spacing = 1.4
    pf.space_after = Pt(5)
    for name, size, before, after in (("Heading 1", 15, 16, 8), ("Heading 2", 12, 12, 5), ("Heading 3", 10.5, 8, 3)):
        _style_font(st[name], SERIF, size, bold=True)
        st[name].paragraph_format.space_before = Pt(before)
        st[name].paragraph_format.space_after = Pt(after)
        st[name].paragraph_format.keep_with_next = True
    _style_font(st["Title"], SERIF, 22, bold=True)
    for lst in ("List Bullet", "List Number"):
        _style_font(st[lst], SANS, 10)
        st[lst].paragraph_format.space_after = Pt(3)
        st[lst].paragraph_format.line_spacing = 1.35
    add_page_number_footer(doc)
    return doc


def add_page_number_footer(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    _font(run, size=8, color=MUTED)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._element.append(el)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run._element.append(instr)


INLINE = re.compile(r"(\*\*.+?\*\*|〔.+?〕)")


def rich(paragraph, text, size=None, color=None, base_bold=False):
    """Supports **bold**. Evidence labels such as 〔计算〕 are kept in the demo page but left out of the report."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            _font(r, size=size, bold=True, color=color or INK)
        elif part.startswith("〔") and part.endswith("〕"):
            continue
        else:
            r = paragraph.add_run(part)
            _font(r, size=size, bold=base_bold, color=color)
    return paragraph


def para(doc, text, size=None, color=None, align=None, space_after=None, bold=False):
    p = doc.add_paragraph()
    rich(p, text, size=size, color=color, base_bold=bold)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def bullets(doc, items, numbered=False):
    for it in items:
        p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
        rich(p, it)


def note(doc, text):
    p = doc.add_paragraph()
    rich(p, text, size=8, color=MUTED)
    p.paragraph_format.space_after = Pt(8)
    return p


def _shade(cell, hex_fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def _cell_borders(table, color="D9DEE6"):
    tbl = table._element
    tblpr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        borders.append(el)
    for edge in ("left", "right", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblpr.append(borders)


def table(doc, header, rows, widths_cm=None, size=8.5, header_fill="1B2433", zebra=False, bold_first_col=False):
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell_borders(t)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        _font(r, size=size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        _shade(c, header_fill)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            rich(p, str(val), size=size, base_bold=(bold_first_col and i == 0))
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.2
            if zebra and ri % 2 == 1:
                _shade(cells[i], "F5F6F8")
    if widths_cm:
        t.autofit = False
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        t._element.tblPr.append(layout)
        grid = t._element.find(qn("w:tblGrid"))
        if grid is not None:
            for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
                gc.set(qn("w:w"), str(int(Cm(w).twips)))
        for row in t.rows:
            for i, w in enumerate(widths_cm):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def callout(doc, title, lines, fill="EEF2FB", title_color=BLUE):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    _shade(c, fill)
    p = c.paragraphs[0]
    r = p.add_run(title)
    _font(r, name=SERIF, size=11, bold=True, color=title_color)
    for ln in lines:
        q = c.add_paragraph()
        rich(q, ln, size=9.5)
        q.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(doc, path, width_cm=16.2, caption=None):
    doc.add_picture(str(path), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        note(doc, caption)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
